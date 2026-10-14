"""1992年風速Excelの5月・6月逆順と、6月の31日空欄を再現する。"""
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import import_hitachi_city_hall_historical_year as historical
import verify_1996_source_cells as oracle

class NamedMonthSelectionTests(unittest.TestCase):
    def setUp(self):
        self.book = Workbook()
        self.book.remove(self.book.active)
        for month, days in ((6, 30), (5, 31)):
            sheet = self.book.create_sheet(f'{month}月')
            sheet['B1'] = '風速'
            sheet['Z1'] = '1992年'
            sheet['AB1'] = month
            sheet['AC1'] = '月'
            for column in range(2,26): sheet.cell(2,column,column-1)
            for day in range(1,32):
                sheet.cell(day+3,1,day)
                if day<=days:
                    for column in range(2,26): sheet.cell(day+3,column,float(month))

    def test_names_headers_and_calendar_days_override_sheet_order(self):
        with patch.object(historical, 'historical_source_workbook', return_value=self.book), \
             patch.object(oracle, 'oracle_source_workbook', return_value=self.book):
            for month in (5,6):
                sheet=historical.month_sheet(Path('wv1992.xlsx'),month)
                self.assertEqual(sheet.title,f'{month}月')
                checked, _, _=oracle.source_table({'wind_speed':Path('wv1992.xlsx')},'wind_speed_m_s',month,1992)
                self.assertIs(checked,sheet)
        self.assertEqual(self.book['5月']['B34'].value,5.0)
        self.assertIsNone(self.book['6月']['B34'].value)

    def test_observation_on_june_31_stops_selection(self):
        self.book['6月']['B34']=1.0
        with patch.object(historical,'historical_source_workbook',return_value=self.book):
            with self.assertRaisesRegex(ValueError,'Observation-day mismatch'):
                historical.month_sheet(Path('wv1992.xlsx'),6)

    def test_conflicting_month_heading_stops_selection(self):
        self.book['5月']['AB1']=6
        with patch.object(historical,'historical_source_workbook',return_value=self.book):
            with self.assertRaisesRegex(ValueError,'Month header mismatch'):
                historical.month_sheet(Path('wv1992.xlsx'),5)
