"""原資料診断38493149507の気圧時刻と集計欄を回帰確認する。"""
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import import_hitachi_city_hall_historical_year as historical
import verify_1996_source_cells as oracle

class EarlierPressureLayoutTests(unittest.TestCase):
    def test_four_daily_pressure_observations_do_not_read_summary_as_hours(self):
        for year in (1986, 1987, 1988, 1989, 1990, 1991, 1992):
            sheet = Workbook().active
            sheet.title = '1月'
            sheet.cell(1, 2, '現地気圧')
            for column, value in enumerate((3, 9, 15, 21, '日平均', '最高', '時刻1', '最低', '時刻2'), 2):
                sheet.cell(2, column, value)
            for day in range(1, 32):
                sheet.cell(day + 2, 1, day)
                for column in range(2, 6): sheet.cell(day + 2, column, 1000 + column)
                sheet.cell(day + 2, 6, 1003)
            with patch.object(historical, 'month_sheet', return_value=sheet):
                extraction = historical.extract_hourly(Path('pressure'), year, 1,
                    'station_pressure_hpa', list(range(1,25)), title_fragment='現地気圧')
            self.assertEqual(len(extraction.cells), 124)
            self.assertEqual({hour for _, hour in extraction.cells}, {3,9,15,21})
            self.assertEqual(oracle.table_specs(year)['station_pressure_hpa'][1], (3,9,15,21))
            self.assertNotIn((1,6), extraction.cells)
