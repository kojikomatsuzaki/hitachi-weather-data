"""古い帳票の時刻と集計位置を、毎時値や新しい帳票に寄せない。"""
import sys
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import import_hitachi_city_hall_historical_year as historical
import verify_1996_source_cells as oracle

class OlderLayoutTests(unittest.TestCase):
    def test_six_hour_temperature_and_humidity_summary_columns(self):
        for element, labels in [('temperature',('日平均','最高','時刻1','最低','時刻2')),('humidity',('日平均','最小','時刻1'))]:
            sheet=Workbook().active;sheet.title='1月'
            for column,value in enumerate((3,9,15,21)+labels,2):sheet.cell(2,column,value)
            for day in range(1,32):
                sheet.cell(day+2,1,day)
                for col in range(2,6):sheet.cell(day+2,col,col)
                sheet.cell(day+2,6,3.5);sheet.cell(day+2,7,7)
                sheet.cell(day+2,8,'09:00')
                if element=='temperature':sheet.cell(day+2,9,-2);sheet.cell(day+2,10,'03:00')
            # The other tables are deliberately omitted from this focused fixture.
            with patch.object(historical,'month_sheet',return_value=sheet):
                with patch.object(historical,'find_title_row',side_effect=RuntimeError('stop after first two tables')):
                    specs=oracle.table_specs(1971)
                    result=historical.extract_hourly(Path('source'),1971,1, 'temperature_c' if element=='temperature' else 'relative_humidity_percent',list(range(1,25)))
                    self.assertEqual(len(result.cells),124)
                    self.assertEqual({h for _,h in result.cells},{3,9,15,21})
                    self.assertEqual(oracle.summary_column(1971,element,10),6)
                    self.assertEqual(specs[result.element_id][1],(3,9,15,21))

    def test_weather_is_serialized_at_nine_without_moving_to_noon(self):
        for year in (1971,1973):
            extractions=[historical.common.ElementExtraction(e,{}) for e in historical.common.HOURLY_VALUE_ORDER]
            extractions.append(historical.common.ElementExtraction('weather_code',{(day,9):historical.common.SourceCell(1) for day in range(1,32)}))
            with patch.object(historical.common,'YEAR',year):d=historical.common.assemble_month_document(1,extractions)
            self.assertEqual(d['observations'][8]['values']['weather_code'],1)
            self.assertNotIn('weather_code',d['observations'][11]['values'])
            self.assertEqual(oracle.table_specs(year)['weather_code'][1],(9,))
        self.assertEqual(oracle.table_specs(1972)['weather_code'][1],(12,))

    def test_multiple_file_absences_are_explicit_and_bounded(self):
        for key in ('dew_point_temperature','solar_radiation','sunshine_duration'):
            e=historical.unavailable_source_element(key,1977,1)
            self.assertEqual(len(e.cells),744)
            self.assertTrue(all(c.flag=='source_file_unavailable' and c.value is None for c in e.cells.values()))
        with self.assertRaises(ValueError):historical.unavailable_source_element('sunshine_duration',1978,1)

    def test_validation_report_uses_original_weather_hour(self):
        for hour in (9,12):
            extractions=[historical.common.ElementExtraction(e,{}) for e in historical.common.HOURLY_VALUE_ORDER]
            extractions.append(historical.common.ElementExtraction('weather_code',{(day,hour):historical.common.SourceCell(1) for day in range(1,32)}))
            with tempfile.TemporaryDirectory() as temp, patch.object(historical.common,'YEAR',1973), patch.object(historical.common,'REPOSITORY_ROOT',Path(temp)):
                d=historical.common.assemble_month_document(1,extractions)
                d['dataset']['observation_schedules']['weather_code']=[hour]
                report=Path(temp)/'report.md'
                historical.common.write_validation_report(report,1,Path(temp)/'data.yaml',d,extractions,[])
                self.assertIn(f'1973-01-01 {hour}時',report.read_text())
