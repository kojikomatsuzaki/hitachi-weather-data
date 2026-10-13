"""1995年診断の列配置で、日付再掲欄を測定値へ混入させない。"""
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import import_hitachi_city_hall_historical_year as historical

class Historical1995SummaryTests(unittest.TestCase):
    def test_dew_point_minimum_skips_repeated_day(self):
        sheets = {}
        for key in ('temperature', 'humidity', 'dew_point_temperature', 'pressure'):
            sheet = Workbook().active
            for header in ([2, 38] if key == 'pressure' else [2]):
                if key == 'pressure': sheet.cell(header-1, 2, '現地気圧' if header == 2 else '海面気圧')
                for col, hour in enumerate(range(3,25,3),2): sheet.cell(header,col,hour)
                labels = {10:'日平均',11:'最高',12:'時刻1',13:'最低',14:'時刻2'}
                if key in ('pressure','dew_point_temperature'): labels={10:'日平均',11:'最高',12:'時刻1',13:'日',14:'最低',15:'時刻2'}
                if key == 'humidity': labels={10:'日平均',11:'最小',12:'時刻1'}
                for col,label in labels.items(): sheet.cell(header,col,label)
                for day in range(1,32):
                    sheet.cell(header+day,1,day)
                    for col in (10,11,13,14): sheet.cell(header+day,col,day if col==13 else -10)
                    sheet.cell(header+day,12,'07:43');sheet.cell(header+day,15,'15:03')
            sheets[key]=sheet
        with patch.object(historical,'month_sheet',side_effect=lambda key,month:sheets[key]):
            result=historical.extract_verified_1996_three_hour_summaries({key:key for key in sheets},1,1995)
        self.assertEqual(result['dew_point_temperature']['1995-01-01']['minimum'],{'value_c':-10.0,'time':'15:03'})
