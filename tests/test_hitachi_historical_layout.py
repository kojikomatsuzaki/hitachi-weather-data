"""Independent fixed Excel-cell expectations from diagnostic run 38092734697."""
import sys
import unittest
from pathlib import Path
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from hitachi_historical_layout import (
    THREE_HOURLY, HOURLY, DAYLIGHT, detect_observation_layout,
    extract_observed_cells,
)


class HistoricalLayoutTests(unittest.TestCase):
    def sheet(self, hours, labels=()):
        wb = Workbook()
        ws = wb.active
        ws.title = "1月"
        for col, hour in enumerate(hours, 2):
            ws.cell(2, col, hour)
        for col, label in enumerate(labels, len(hours) + 2):
            ws.cell(2, col, label)
        for day in range(1, 32):
            ws.cell(day + 2, 1, day)
        return ws

    def test_three_hour_header_and_original_january_values(self):
        # Source: 1996 temperature.xls, 1月, B2:N3 (diagnostic Artifact).
        ws = self.sheet(THREE_HOURLY, ("日平均", "最高", "時刻1", "最低", "時刻2"))
        originals = (6.5, 4.9, 8.2, 9.9, 8.5, 3.7, 1.2, 0.3)
        for col, value in enumerate(originals, 2):
            ws.cell(3, col, value)
        for col, value in enumerate((5.4, 10.9, "11:47", -0.7, "23:22"), 10):
            ws.cell(3, col, value)
        layout = detect_observation_layout(
            ws, expected_schedules=(THREE_HOURLY,),
            summary_labels=("日平均", "最高", "時刻1", "最低", "時刻2"))
        cells = extract_observed_cells(ws, layout, 1996, 1, lambda value: value)
        self.assertEqual([cells[(1, hour)] for hour in THREE_HOURLY], list(originals))
        self.assertNotIn((1, 1), cells)
        self.assertNotIn((1, 2), cells)
        self.assertEqual(ws["J3"].value, 5.4)
        self.assertEqual(ws["K3"].value, 10.9)
        self.assertEqual(ws["M3"].value, -0.7)
        self.assertEqual(len(cells), 31 * 8)

    def test_hourly_regression(self):
        ws = self.sheet(HOURLY)
        layout = detect_observation_layout(ws, expected_schedules=(HOURLY,))
        self.assertEqual(layout.hours, HOURLY)
        self.assertEqual(layout.columns, tuple(range(2, 26)))

    def test_daylight_header(self):
        ws = self.sheet(DAYLIGHT)
        layout = detect_observation_layout(ws, expected_schedules=(DAYLIGHT,))
        self.assertEqual(layout.hours, DAYLIGHT)

    def test_unknown_header_rejected(self):
        ws = self.sheet((3, 6, 9, 13, 15, 18, 21, 24))
        with self.assertRaises(ValueError):
            detect_observation_layout(ws)

    def test_summary_mismatch_rejected(self):
        ws = self.sheet(THREE_HOURLY, ("日平均", "最高", "不明", "最低", "時刻2"))
        with self.assertRaisesRegex(ValueError, "Summary header mismatch"):
            detect_observation_layout(
                ws, expected_schedules=(THREE_HOURLY,),
                summary_labels=("日平均", "最高", "時刻1", "最低", "時刻2"))

    def test_leap_day(self):
        ws = self.sheet(THREE_HOURLY)
        for day in range(1, 32):
            ws.cell(day + 2, 1, day if day <= 29 else None)
        layout = detect_observation_layout(ws, expected_schedules=(THREE_HOURLY,))
        cells = extract_observed_cells(ws, layout, 1996, 2, lambda value: value)
        self.assertIn((29, 24), cells)
        self.assertNotIn((30, 24), cells)

    def test_blank_not_interpolated(self):
        ws = self.sheet(THREE_HOURLY)
        layout = detect_observation_layout(ws, expected_schedules=(THREE_HOURLY,))
        cells = extract_observed_cells(ws, layout, 1996, 1, lambda value: value)
        self.assertIsNone(cells[(1, 3)])
        self.assertNotIn((1, 4), cells)


if __name__ == "__main__":
    unittest.main()
