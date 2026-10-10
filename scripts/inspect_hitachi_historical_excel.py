#!/usr/bin/env python3
"""Read-only inventory of historical Excel ZIP members and worksheet layouts.

The original ZIP and XLS bytes are never modified. Conversion happens in a
temporary directory. This script intentionally does not infer observation values.
"""
import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "metadata/sources/hitachi-city-hall-historical-archives.yaml"


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def cell_value(value):
    if value is None:
        return None
    return str(value)


def inspect_member(member_bytes, member_name, year, element, archive_name):
    with tempfile.TemporaryDirectory(prefix="hitachi-inspect-") as temporary:
        directory = Path(temporary)
        source = directory / Path(member_name).name
        source.write_bytes(member_bytes)
        converted = directory / "converted"
        converted.mkdir()
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if office is None:
            raise RuntimeError("LibreOffice is required")
        result = subprocess.run(
            [office, "--headless",
             f"-env:UserInstallation={(directory / 'profile').as_uri()}",
             "--convert-to", "xlsx", "--outdir", str(converted), str(source)],
            capture_output=True, text=True, timeout=180, check=True)
        output = converted / (source.stem + ".xlsx")
        if not output.exists():
            raise RuntimeError("Conversion produced no XLSX: " + result.stdout)
        workbook = load_workbook(output, data_only=True, read_only=False)
        sheets = []
        for sheet in workbook.worksheets:
            preview = []
            candidates = []
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 45),
                                       max_col=min(sheet.max_column, 45)):
                nonempty = [{"cell": cell.coordinate, "value": cell_value(cell.value)}
                            for cell in row if cell.value is not None]
                if nonempty:
                    preview.append({"row": row[0].row, "cells": nonempty})
                for index in range(len(row) - 23):
                    values = [row[index + offset].value for offset in range(24)]
                    def integer(value):
                        try:
                            number = float(value)
                            return int(number) if number.is_integer() else None
                        except (ValueError, TypeError):
                            return None
                    if [integer(v) for v in values] == list(range(1, 25)):
                        candidates.append({"first_cell": row[index].coordinate,
                                           "last_cell": row[index + 23].coordinate,
                                           "kind": "1-24"})
            sheets.append({
                "name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                "preview": preview, "hour_header_candidates": candidates,
                "column_mapping_status": "not_verified",
                "daily_summary_mapping_status": "not_verified"
            })
        return {"year": year, "element": element, "archive": archive_name,
                "excel": member_name, "excel_sha256": sha256(member_bytes),
                "excel_format": "legacy XLS (converted read-only via LibreOffice)",
                "sheet_count": len(sheets), "sheets": sheets}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--years", nargs="+", type=int, default=[1996, 1997])
    parser.add_argument("--cache-directory", type=Path,
                        default=ROOT / "build/historical-archives")
    parser.add_argument("--output-directory", type=Path,
                        default=ROOT / "build/historical-inspection")
    args = parser.parse_args()
    args.output_directory.mkdir(parents=True, exist_ok=True)
    catalog = yaml.safe_load(CATALOG.read_text(encoding="utf-8"))
    import urllib.request
    inventory, structures, failures = [], [], []
    for archive in catalog["archives"]:
        path = args.cache_directory / archive["file_name"]
        origin = "cache" if path.is_file() else "download"
        if not path.is_file():
            try:
                path.parent.mkdir(parents=True, exist_ok=True)
                urllib.request.urlretrieve(archive["url"], path)
            except Exception as error:
                failures.append({"archive": archive["file_name"], "status": "download_failed",
                                 "error": str(error)})
                continue
        content = path.read_bytes()
        checksum = sha256(content)
        inventory.append({"archive": path.name, "origin": origin,
                          "bytes": len(content), "sha256": checksum,
                          "expected_sha256": archive["sha256"],
                          "verified": checksum == archive["sha256"]})
        if checksum != archive["sha256"]:
            failures.append({"archive": path.name, "status": "checksum_mismatch"})
            continue
        with zipfile.ZipFile(path) as zip_file:
            members = zip_file.namelist()
            inventory[-1]["members"] = members
            for year in args.years:
                target = (archive["member_prefix"] + str(year) + ".xls").casefold()
                matches = [m for m in members if Path(m).name.casefold() == target]
                if len(matches) != 1:
                    failures.append({"archive": path.name, "year": year,
                                     "status": "member_missing_or_ambiguous",
                                     "matches": matches})
                    continue
                try:
                    structure = inspect_member(zip_file.read(matches[0]), matches[0],
                                               year, archive["element"], path.name)
                    structures.append(structure)
                    print(f"INSPECTED {year} {archive['element']} {matches[0]}", flush=True)
                    for sheet in structure["sheets"]:
                        if sheet["name"] in ("1月", "2月"):
                            print(f"  {sheet['name']} headers={sheet['hour_header_candidates']}", flush=True)
                            for row in sheet["preview"][:8]:
                                print(f"  row {row['row']}: {row['cells'][:32]}", flush=True)
                except Exception as error:
                    failures.append({"archive": path.name, "year": year,
                                     "status": "inspection_failed", "error": str(error)})
    (args.output_directory / "source-checksums.json").write_text(
        json.dumps({"archives": inventory, "failures": failures},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output_directory / "excel-structure.json").write_text(
        json.dumps(structures, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.output_directory / "excel-structure.csv").open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=[
            "year", "element", "archive", "excel", "sheet", "rows", "columns",
            "hour_header_candidates", "merged_ranges"])
        writer.writeheader()
        for book in structures:
            for sheet in book["sheets"]:
                writer.writerow({"year": book["year"], "element": book["element"],
                                 "archive": book["archive"], "excel": book["excel"],
                                 "sheet": sheet["name"], "rows": sheet["rows"],
                                 "columns": sheet["columns"],
                                 "hour_header_candidates": json.dumps(sheet["hour_header_candidates"]),
                                 "merged_ranges": json.dumps(sheet["merged_ranges"])})
    lines = ["# Historical Excel diagnostic", "",
             f"Inspected workbooks: {len(structures)}",
             f"Failures: {len(failures)}", "",
             "Column semantics and daily summaries are NOT automatically verified.",
             "Inspect cell previews before changing the importer.", ""]
    for failure in failures:
        lines.append("- " + json.dumps(failure, ensure_ascii=False))
    (args.output_directory / "diagnostic-summary.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8")
    pairs = {(s["year"], s["element"]): s for s in structures}
    comparisons = ["# 1996 vs 1997", "", "Preliminary structural comparison only.", ""]
    for element in [a["element"] for a in catalog["archives"]]:
        a, b = pairs.get((1996, element)), pairs.get((1997, element))
        comparisons.append(f"## {element}")
        if not a or not b:
            comparisons.append("One or both workbooks unavailable.")
            continue
        for index in range(min(len(a["sheets"]), len(b["sheets"]))):
            x, y = a["sheets"][index], b["sheets"][index]
            comparisons.append(f"- {x['name']} / {y['name']}: "
                               f"dimensions {x['rows']}x{x['columns']} / {y['rows']}x{y['columns']}; "
                               f"hour candidates {x['hour_header_candidates']} / {y['hour_header_candidates']}")
    (args.output_directory / "comparison-1996-1997.md").write_text(
        "\n".join(comparisons) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
