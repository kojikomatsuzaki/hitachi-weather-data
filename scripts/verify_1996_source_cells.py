"""Independent 1996 worksheet oracle; never calls production layout/decoders.

Reads checksum-verified XLS originals via their LibreOffice XLSX derivatives.
Coordinates and schedules are the source table definitions, not inferred from
production output. Does not certify the LibreOffice conversion itself.
"""
import calendar
import math
import unicodedata
from collections import Counter
from datetime import datetime, time, timedelta

from openpyxl import load_workbook

THREE = (3, 6, 9, 12, 15, 18, 21, 24)
FULL = tuple(range(1, 25))
LIGHT = tuple(range(4, 21))
# source key, schedule, pressure table title
TABLES = {
    'temperature_c': ('temperature', THREE, None),
    'relative_humidity_percent': ('humidity', THREE, None),
    'dew_point_temperature_c': ('dew_point_temperature', THREE, None),
    'station_pressure_hpa': ('pressure', THREE, '現地気圧'),
    'sea_level_pressure_hpa': ('pressure', THREE, '海面気圧'),
    'precipitation_mm': ('precipitation', FULL, None),
    'wind_speed_m_s': ('wind_speed', FULL, None),
    'wind_direction': ('wind_direction', FULL, None),
    'global_solar_radiation_mj_m2': ('solar_radiation', LIGHT, None),
    'sunshine_duration_h': ('sunshine_duration', LIGHT, None),
    'weather_code': ('weather_code', (12,), None),
}
DIRECTIONS = dict(zip(
    ('北 北北東 北東 東北東 東 東南東 南東 南南東 南 南南西 南西 西南西 西 西北西 北西 北北西 静穏').split(),
    ('N NNE NE ENE E ESE SE SSE S SSW SW WSW W WNW NW NNW CALM').split()))
CODES = {0, 1, 2, 3, 10, 40, 50, 60, 70, 80, 85, 88, 90}
SPECIAL = {'-': 'source_dash', '***': 'source_triple_asterisk', '欠測': 'source_missing'}


def decode(raw, element):
    """Return (value, flag, raw_value, has_raw_value) independently."""
    if raw is None or raw == '' or isinstance(raw, float) and math.isnan(raw):
        return None, 'source_blank', None, False
    if isinstance(raw, str):
        stripped = raw.strip()
        if not stripped:
            return None, 'source_whitespace', raw, True
        if stripped in SPECIAL:
            return None, SPECIAL[stripped], raw, True
        if element == 'wind_direction' and stripped in DIRECTIONS.values():
            return stripped, None, None, False
        if element == 'wind_direction' and stripped in DIRECTIONS:
            return DIRECTIONS[stripped], None, None, False
        return None, 'unrecognized_source_value', raw, True
    if element == 'sunshine_duration_h' and isinstance(raw, timedelta):
        return raw.total_seconds() / 3600, None, None, False
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        if element == 'wind_direction':
            return None, 'unrecognized_source_value', float(raw), True
        if element == 'weather_code':
            # No truncation of fractional codes allowed.
            if raw != int(raw) or int(raw) not in CODES:
                return None, 'unrecognized_source_value', raw, True
            return int(raw), None, None, False
        return float(raw), None, None, False
    raise ValueError(f'Unsupported source value {element}: {raw!r}')


def number(sheet, row, col, duration=False):
    return decode(sheet.cell(row, col).value,
                  'sunshine_duration_h' if duration else 'numeric')[0]


def moment(sheet, row, col):
    raw = sheet.cell(row, col).value
    if raw is None or raw == '':
        return None
    if isinstance(raw, datetime):
        raw = raw.time()
    if isinstance(raw, time):
        return f'{raw.hour:02}:{raw.minute:02}'
    if isinstance(raw, timedelta):
        seconds = int(raw.total_seconds())
        hours, seconds = divmod(seconds, 3600)
        minutes, seconds = divmod(seconds, 60)
        return f'{hours:02}:{minutes:02}' + (f':{seconds:02}' if seconds else '')
    return str(raw)


def text_cell(sheet, row, col):
    raw = sheet.cell(row, col).value
    if raw is None or raw == '' or isinstance(raw, str) and (
            not raw.strip() or raw.strip() in SPECIAL):
        return None
    return raw.strip() if isinstance(raw, str) else str(float(raw))


def source_table(paths, element, month):
    key, hours, fragment = TABLES[element]
    workbook = load_workbook(paths[key], data_only=True)
    if element == 'weather_code':
        sheet = workbook['天気']
        assert '12時' in unicodedata.normalize('NFKC', str(sheet.cell(1, 1).value))
        assert str(sheet.cell(3, 19 + month).value).strip() == str(month)
        return sheet, 4, (19 + month,)
    sheet = workbook.worksheets[month - 1]
    assert unicodedata.normalize('NFKC', sheet.title) == f'{month}月'
    start = 1
    if fragment:
        title_rows = [row for row in range(1, sheet.max_row + 1)
                      if any(fragment in str(sheet.cell(row, col).value or '')
                             for col in range(1, 9))]
        assert len(title_rows) == 1, (element, month, title_rows)
        start = title_rows[0] + 1
    headers = [row for row in range(start, min(start + 8, sheet.max_row) + 1)
               if tuple(sheet.cell(row, col).value for col in range(2, 2 + len(hours))) == hours]
    assert len(headers) == 1, (element, month, headers)
    return sheet, headers[0], tuple(range(2, 2 + len(hours)))


def expected_cells(paths, month):
    output, addresses = {}, {}
    days = calendar.monthrange(1996, month)[1]
    for element, (_, hours, _) in TABLES.items():
        sheet, header, columns = source_table(paths, element, month)
        rows = {}
        for row in range(header + 1, sheet.max_row + 1):
            raw_day = sheet.cell(row, 1).value
            if isinstance(raw_day, (int, float)) and raw_day == int(raw_day) and 1 <= raw_day <= days:
                day = int(raw_day)
                if day in rows:
                    break  # next table, e.g. sea-level pressure
                rows[day] = row
                if len(rows) == days:
                    break
        assert set(rows) == set(range(1, days + 1)), (element, month, rows)
        cells, refs = {}, {}
        for day, row in rows.items():
            for hour, col in zip(hours, columns):
                cells[day, hour] = decode(sheet.cell(row, col).value, element)
                refs[day, hour] = f'{sheet.title}!{sheet.cell(row, col).coordinate}'
        output[element], addresses[element] = cells, refs
    return output, addresses


def verify_extractions(paths, month, extractions):
    expected, refs = expected_cells(paths, month)
    assert {item.element_id for item in extractions} == set(TABLES)
    result = {}
    for item in extractions:
        oracle = expected[item.element_id]
        assert set(item.cells) == set(oracle), (month, item.element_id, 'schedule')
        for key, (value, flag, raw, has_raw) in oracle.items():
            actual = item.cells[key]
            assert (actual.value, actual.flag) == (value, flag), (
                month, item.element_id, key, refs[item.element_id][key], actual, oracle[key])
            # Known wind tokens normalize without a raw_values entry.
            assert actual.has_raw_value == has_raw
            if has_raw:
                assert actual.raw_value == raw
        result[item.element_id] = {
            'hours': list(TABLES[item.element_id][1]),
            'observation_cells': len(oracle),
            'values_present': sum(cell[0] is not None for cell in oracle.values()),
            'flags': dict(Counter(cell[1] for cell in oracle.values() if cell[1])),
            'independent_source_comparison': True,
        }
    return result


def verify_document(paths, month, document):
    expected, refs = expected_cells(paths, month)
    days = calendar.monthrange(1996, month)[1]
    assert document['dataset']['year'] == 1996 and document['dataset']['month'] == month
    assert document['dataset']['observation_schedules'] == {
        element: list(spec[1]) for element, spec in TABLES.items()}
    records = {(int(record['source_date'][-2:]), record['source_hour']): record
               for record in document['observations']}
    assert len(records) == len(document['observations']) == days * 24
    assert set(records) == {(day, hour) for day in range(1, days + 1) for hour in FULL}
    for (day, hour), record in records.items():
        assert record['source_date'] == f'1996-{month:02}-{day:02}'
        values, flags, raw_values = {}, {}, {}
        for element, cells in expected.items():
            if (day, hour) in cells:
                value, flag, raw, has_raw = cells[day, hour]
                values[element] = value
                if flag:
                    flags[element] = flag
                if flag and has_raw:
                    raw_values[element] = raw
            elif element != 'weather_code':
                values[element] = None
                flags[element] = 'not_scheduled'
        assert record['values'] == values, (month, day, hour, 'values')
        assert record.get('flags', {}) == flags, (month, day, hour, 'flags')
        assert record.get('raw_values', {}) == raw_values, (month, day, hour, 'raw_values')
    return len(records)
