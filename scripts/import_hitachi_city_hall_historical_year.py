#!/usr/bin/env python3
"""歴史資料期の要素別ZIPから、指定年の月別YAMLを生成する。

古いExcelには、xlrdが読み取れないBIFFレコードを含むものがある。
原Excel自体は変更せず、LibreOffice Calcで一時XLSXへ変換し、その値を読む。
ZIPと内部Excelの双方をSHA-256で検証してから変換する。
"""

from __future__ import annotations

import argparse
import calendar
import hashlib
import shutil
import subprocess
import sys
import urllib.request
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

import import_hitachi_city_hall_2025 as common


# ==========================================
# Repository paths and source definitions
# ==========================================

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
STATION_ID = "hitachi-city-hall"
GENERATOR_PATH = "scripts/import_hitachi_city_hall_historical_year.py"

ARCHIVE_FILE_NAMES = {
    "temperature": "temperature.zip",
    "humidity": "humidity.zip",
    "precipitation": "precipitation.zip",
    "pressure": "pressure.zip",
    "solar_radiation": "global_solar_radiation.zip",
    "sunshine_duration": "sunshine_duration.zip",
    "wind_speed": "wind_speed.zip",
    "wind_direction": "wind_direction.zip",
    "dew_point_temperature": "dew_point_temperature.zip",
    "weather_code": "weather_code.zip",
}


# ==========================================
# Source acquisition and preservation
# ==========================================

def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source_file:
        for block in iter(lambda: source_file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_manifest(year: int) -> tuple[Path, dict[str, Any]]:
    path = REPOSITORY_ROOT / "metadata" / "sources" / f"hitachi-city-hall-{year}.yaml"
    return path, yaml.safe_load(path.read_text(encoding="utf-8"))


def convert_workbook_without_changing_source(source_path: Path, output_directory: Path) -> Path:
    """原XLSを不変のまま保ち、読取専用の一時XLSXを作る。"""

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise FileNotFoundError(
            "LibreOffice Calc is required to read the historical Excel files."
        )

    output_directory.mkdir(parents=True, exist_ok=True)
    output_path = output_directory / f"{source_path.stem}.xlsx"
    # build配下の変換済みファイルは同じ原Excelから再生成できるキャッシュである。
    # 月ごとの実行で10冊を毎回変換し直さない。
    if output_path.exists():
        return output_path
    output_path.unlink(missing_ok=True)
    profile_directory = output_directory / f"libreoffice-profile-{source_path.stem}"
    profile_directory.mkdir(parents=True, exist_ok=True)

    subprocess.run(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation={profile_directory.resolve().as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_directory),
            str(source_path),
        ],
        check=True,
        timeout=180,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    if not output_path.exists():
        raise ValueError(f"LibreOffice did not create {output_path.name}")
    return output_path


def prepare_source_files(
    manifest: dict[str, Any],
    source_directory: Path,
    download_missing: bool,
) -> tuple[dict[str, Path], list[dict[str, str]], list[dict[str, str]]]:
    """ZIP・内部Excelを照合し、読取用XLSXを準備する。"""

    original_directory = source_directory / "original"
    converted_directory = source_directory / "converted"
    original_directory.mkdir(parents=True, exist_ok=True)

    source_paths: dict[str, Path] = {}
    workbook_rows: list[dict[str, str]] = []
    archive_rows: list[dict[str, str]] = []

    for entry in manifest["files"]:
        element = entry["element"]
        archive_path = source_directory / ARCHIVE_FILE_NAMES[element]
        if not archive_path.exists():
            if not download_missing:
                raise FileNotFoundError(
                    f"Source archive is missing: {archive_path}. Use --download-missing."
                )
            urllib.request.urlretrieve(entry["url"], archive_path)

        archive_digest = sha256_path(archive_path)
        if archive_digest != entry["archive_sha256"]:
            raise ValueError(
                f"SHA-256 mismatch for {archive_path.name}: "
                f"expected {entry['archive_sha256']}, got {archive_digest}"
            )

        with zipfile.ZipFile(archive_path) as archive:
            workbook_bytes = archive.read(entry["archive_member"])
        member_digest = sha256_bytes(workbook_bytes)
        if member_digest != entry["member_sha256"]:
            raise ValueError(
                f"SHA-256 mismatch for {entry['archive_member']}: "
                f"expected {entry['member_sha256']}, got {member_digest}"
            )

        original_path = original_directory / entry["archive_member"]
        if not original_path.exists() or sha256_path(original_path) != member_digest:
            original_path.write_bytes(workbook_bytes)
        converted_path = convert_workbook_without_changing_source(
            original_path, converted_directory
        )
        source_paths[element] = converted_path

        archive_rows.append({
            "element": element,
            "file": archive_path.name,
            "sha256": archive_digest,
            "status": "verified",
        })
        workbook_rows.append({
            "element": element,
            "file": entry["archive_member"],
            "sha256": member_digest,
            "status": "verified",
        })

    return source_paths, workbook_rows, archive_rows


# ==========================================
# Worksheet helpers
# ==========================================

def month_sheet(path: Path, month: int):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.worksheets) < 12:
        raise ValueError(f"Expected 12 monthly sheets in {path.name}")
    return workbook.worksheets[month - 1]


def integer_value(value: Any) -> int | None:
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def find_title_row(sheet, fragment: str) -> int:
    for row in range(1, sheet.max_row + 1):
        for column in range(1, min(sheet.max_column, 8) + 1):
            if fragment in str(sheet.cell(row, column).value or ""):
                return row
    raise ValueError(f"Could not find title {fragment!r} in sheet {sheet.title!r}")


def find_hour_header_row(sheet, hours: list[int], start_row: int = 1) -> int:
    for row in range(start_row, min(sheet.max_row, start_row + 8) + 1):
        actual = [integer_value(sheet.cell(row, offset + 2).value) for offset in range(len(hours))]
        if actual == hours:
            return row
    raise ValueError(
        f"Could not find hour header {hours[0]}-{hours[-1]} in {sheet.title!r}"
    )


def day_rows(sheet, header_row: int, year: int, month: int):
    expected_days = calendar.monthrange(year, month)[1]
    found: set[int] = set()
    for row in range(header_row + 1, sheet.max_row + 1):
        day = integer_value(sheet.cell(row, 1).value)
        if day is None or not 1 <= day <= expected_days or day in found:
            continue
        found.add(day)
        yield day, row
        if len(found) == expected_days:
            return
    raise ValueError(f"Expected {expected_days} daily rows; found {len(found)}")


# ==========================================
# Observation extraction
# ==========================================

def extract_hourly(
    path: Path,
    year: int,
    month: int,
    element_id: str,
    hours: list[int],
    *,
    allow_text: bool = False,
    title_fragment: str | None = None,
) -> common.ElementExtraction:
    sheet = month_sheet(path, month)
    start_row = find_title_row(sheet, title_fragment) + 1 if title_fragment else 1
    header_row = find_hour_header_row(sheet, hours, start_row)
    cells: dict[tuple[int, int], common.SourceCell] = {}
    for day, row in day_rows(sheet, header_row, year, month):
        for offset, hour in enumerate(hours, start=2):
            raw_value = sheet.cell(row, offset).value
            if allow_text:
                cell = common.source_cell(raw_value)
            else:
                cell = historical_numeric_cell(
                    raw_value,
                    duration_unit_hours=element_id == "sunshine_duration_h",
                )
            cells[(day, hour)] = cell
    return common.ElementExtraction(element_id, cells)


def historical_numeric_cell(
    value: Any,
    *,
    duration_unit_hours: bool = False,
) -> common.SourceCell:
    """Excelの時間長を、時間単位の観測値として明示的に変換する。

    openpyxlは時間長のセルを`timedelta`として返すことがある。
    日照時間の帳票に限り、秒から時間へ換算する。別の項目に現れた
    `timedelta`は意味を決めず、既存の厳格な検証で停止させる。
    """

    if duration_unit_hours and isinstance(value, timedelta):
        return common.SourceCell(value.total_seconds() / 3600)
    return common.numeric_source_cell(value)


def extract_wind_direction(path: Path, year: int, month: int) -> common.ElementExtraction:
    raw = extract_hourly(
        path, year, month, "wind_direction", list(range(1, 25)), allow_text=True
    )
    converted: dict[tuple[int, int], common.SourceCell] = {}
    for key, cell in raw.cells.items():
        if cell.value is None:
            converted[key] = cell
        elif cell.value in common.WIND_DIRECTION_CODES:
            converted[key] = common.SourceCell(common.WIND_DIRECTION_CODES[cell.value], cell.flag)
        else:
            raw_value = cell.raw_value if cell.has_raw_value else cell.value
            converted[key] = common.SourceCell(
                None, "unrecognized_source_value", raw_value, True
            )
    return common.ElementExtraction("wind_direction", converted)


def extract_weather(path: Path, year: int, month: int) -> common.ElementExtraction:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "天気" not in workbook.sheetnames:
        raise ValueError(f"Weather sheet was not found in {path.name}")
    sheet = workbook["天気"]
    code_column = 19 + month
    if integer_value(sheet.cell(3, code_column).value) != month:
        raise ValueError(f"Could not locate weather code column for month {month}")
    cells: dict[tuple[int, int], common.SourceCell] = {}
    for day, row in day_rows(sheet, 4, year, month):
        cell = common.source_cell(sheet.cell(row, code_column).value, integer=True)
        if cell.value is not None and cell.value not in common.WEATHER_CODES:
            raw_value = cell.raw_value if cell.has_raw_value else cell.value
            cell = common.SourceCell(None, "unrecognized_source_value", raw_value, True)
        cells[(day, 12)] = cell
    return common.ElementExtraction("weather_code", cells)


def extract_all_elements(
    source_paths: dict[str, Path], year: int, month: int
) -> list[common.ElementExtraction]:
    full_day = list(range(1, 25))
    daylight = list(range(4, 21))
    return [
        extract_hourly(source_paths["temperature"], year, month, "temperature_c", full_day),
        extract_hourly(source_paths["humidity"], year, month, "relative_humidity_percent", full_day),
        extract_hourly(source_paths["precipitation"], year, month, "precipitation_mm", full_day),
        extract_hourly(source_paths["pressure"], year, month, "station_pressure_hpa", full_day, title_fragment="現地気圧"),
        extract_hourly(source_paths["pressure"], year, month, "sea_level_pressure_hpa", full_day, title_fragment="海面気圧"),
        extract_hourly(source_paths["solar_radiation"], year, month, "global_solar_radiation_mj_m2", daylight),
        extract_hourly(source_paths["sunshine_duration"], year, month, "sunshine_duration_h", daylight),
        extract_hourly(source_paths["wind_speed"], year, month, "wind_speed_m_s", full_day),
        extract_wind_direction(source_paths["wind_direction"], year, month),
        extract_hourly(source_paths["dew_point_temperature"], year, month, "dew_point_temperature_c", full_day),
        extract_weather(source_paths["weather_code"], year, month),
    ]


# ==========================================
# Daily summaries
# ==========================================

def numeric_value(
    sheet,
    row: int,
    column: int,
    *,
    duration_unit_hours: bool = False,
) -> float | None:
    cell = historical_numeric_cell(
        sheet.cell(row, column).value,
        duration_unit_hours=duration_unit_hours,
    )
    return cell.value if isinstance(cell.value, (int, float)) else None


def text_value(sheet, row: int, column: int) -> str | None:
    cell = common.source_cell(sheet.cell(row, column).value)
    return str(cell.value) if cell.value is not None else None


def time_value(sheet, row: int, column: int) -> str | None:
    value = sheet.cell(row, column).value
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value.isoformat(timespec="minutes")
    if isinstance(value, datetime):
        return value.time().isoformat(timespec="minutes")
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        if total_seconds == 86400:
            return "24:00"
        hour, remainder = divmod(total_seconds, 3600)
        minute, second = divmod(remainder, 60)
        return f"{hour:02d}:{minute:02d}" if second == 0 else f"{hour:02d}:{minute:02d}:{second:02d}"
    return str(value)


def rows_by_day(sheet, header_row: int, year: int, month: int) -> dict[int, int]:
    return {day: row for day, row in day_rows(sheet, header_row, year, month)}


def extract_daily_summaries(
    source_paths: dict[str, Path], year: int, month: int
) -> list[dict[str, Any]]:
    """時間値と混在させず、原帳票の日別集計欄を採録する。"""

    sheets = {
        element: month_sheet(path, month)
        for element, path in source_paths.items()
        if element != "weather_code"
    }
    full_day = list(range(1, 25))
    daylight = list(range(4, 21))
    header_rows = {
        "temperature": find_hour_header_row(sheets["temperature"], full_day),
        "humidity": find_hour_header_row(sheets["humidity"], full_day),
        "precipitation": find_hour_header_row(sheets["precipitation"], full_day),
        "station_pressure": find_hour_header_row(
            sheets["pressure"], full_day, find_title_row(sheets["pressure"], "現地気圧") + 1
        ),
        "sea_level_pressure": find_hour_header_row(
            sheets["pressure"], full_day, find_title_row(sheets["pressure"], "海面気圧") + 1
        ),
        "solar_radiation": find_hour_header_row(sheets["solar_radiation"], daylight),
        "sunshine_duration": find_hour_header_row(sheets["sunshine_duration"], daylight),
        "wind_speed": find_hour_header_row(sheets["wind_speed"], full_day),
        "dew_point_temperature": find_hour_header_row(sheets["dew_point_temperature"], full_day),
    }
    row_maps = {
        key: rows_by_day(
            sheets["pressure"] if key in {"station_pressure", "sea_level_pressure"} else sheets[key],
            header_row,
            year,
            month,
        )
        for key, header_row in header_rows.items()
    }

    summaries: list[dict[str, Any]] = []
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        temperature = sheets["temperature"]
        humidity = sheets["humidity"]
        precipitation = sheets["precipitation"]
        pressure = sheets["pressure"]
        solar = sheets["solar_radiation"]
        sunshine = sheets["sunshine_duration"]
        wind = sheets["wind_speed"]
        dew_point = sheets["dew_point_temperature"]
        rows = {key: mapping[day] for key, mapping in row_maps.items()}

        summaries.append({
            "date": date(year, month, day).isoformat(),
            "temperature": {
                "mean_c": numeric_value(temperature, rows["temperature"], 26),
                "maximum": {
                    "value_c": numeric_value(temperature, rows["temperature"], 27),
                    "time": time_value(temperature, rows["temperature"], 28),
                },
                "minimum": {
                    "value_c": numeric_value(temperature, rows["temperature"], 30),
                    "time": time_value(temperature, rows["temperature"], 31),
                },
            },
            "humidity": {
                "mean_percent": numeric_value(humidity, rows["humidity"], 26),
                "minimum": {
                    "value_percent": numeric_value(humidity, rows["humidity"], 27),
                    "time": time_value(humidity, rows["humidity"], 28),
                },
            },
            "precipitation": {
                "total_mm": numeric_value(precipitation, rows["precipitation"], 26),
                "maximum_one_hour": {
                    "value_mm": numeric_value(precipitation, rows["precipitation"], 28),
                    "time": time_value(precipitation, rows["precipitation"], 29),
                },
                "maximum_ten_minutes": {
                    "value_mm": numeric_value(precipitation, rows["precipitation"], 31),
                    "time": time_value(precipitation, rows["precipitation"], 32),
                },
            },
            "station_pressure": {
                "mean_hpa": numeric_value(pressure, rows["station_pressure"], 26),
                "maximum": {
                    "value_hpa": numeric_value(pressure, rows["station_pressure"], 27),
                    "time": time_value(pressure, rows["station_pressure"], 28),
                },
                "minimum": {
                    "value_hpa": numeric_value(pressure, rows["station_pressure"], 30),
                    "time": time_value(pressure, rows["station_pressure"], 31),
                },
            },
            "sea_level_pressure": {
                "mean_hpa": numeric_value(pressure, rows["sea_level_pressure"], 26),
                "maximum": {
                    "value_hpa": numeric_value(pressure, rows["sea_level_pressure"], 27),
                    "time": time_value(pressure, rows["sea_level_pressure"], 28),
                },
                "minimum": {
                    "value_hpa": numeric_value(pressure, rows["sea_level_pressure"], 30),
                    "time": time_value(pressure, rows["sea_level_pressure"], 31),
                },
            },
            "global_solar_radiation": {
                "total_mj_m2": numeric_value(solar, rows["solar_radiation"], 19),
            },
            "sunshine_duration": {
                "total_h": numeric_value(
                    sunshine,
                    rows["sunshine_duration"],
                    19,
                    duration_unit_hours=True,
                ),
            },
            "wind": {
                "mean_speed_m_s": numeric_value(wind, rows["wind_speed"], 26),
                "maximum_10_minute": {
                    "direction": text_value(wind, rows["wind_speed"], 27),
                    "speed_m_s": numeric_value(wind, rows["wind_speed"], 28),
                    "time": time_value(wind, rows["wind_speed"], 29),
                },
                "maximum_instantaneous": {
                    "direction": text_value(wind, rows["wind_speed"], 31),
                    "speed_m_s": numeric_value(wind, rows["wind_speed"], 32),
                    "time": time_value(wind, rows["wind_speed"], 33),
                },
            },
            "dew_point_temperature": {
                "mean_c": numeric_value(dew_point, rows["dew_point_temperature"], 26),
                "maximum": {
                    "value_c": numeric_value(dew_point, rows["dew_point_temperature"], 27),
                    "time": time_value(dew_point, rows["dew_point_temperature"], 28),
                },
                "minimum": {
                    "value_c": numeric_value(dew_point, rows["dew_point_temperature"], 30),
                    "time": time_value(dew_point, rows["dew_point_temperature"], 31),
                },
            },
        })
    return summaries


# ==========================================
# Canonical document and validation
# ==========================================

def configure_common_module(year: int, manifest_path: Path) -> None:
    common.YEAR = year
    common.SOURCE_MANIFEST_PATH = manifest_path
    common.GENERATOR_PATH = GENERATOR_PATH
    common.SOURCE_PERIOD_ID = "historical_annual_workbooks"


def append_historical_report_details(
    report_path: Path,
    archive_rows: list[dict[str, str]],
) -> None:
    report_text = report_path.read_text(encoding="utf-8").replace(
        "- 日別集計値はこの段階では収録せず、時間観測値とは別の工程で扱う。",
        "- 日別集計値は時間観測値と混在させず、`daily_summaries`へ収録した。",
    )
    report_path.write_text(report_text, encoding="utf-8")
    with report_path.open("a", encoding="utf-8") as report:
        report.write("\n## 原ZIPの整合性確認\n\n")
        report.write("| 要素 | ZIP | SHA-256 | 判定 |\n|---|---|---|---|\n")
        for row in archive_rows:
            report.write(
                f"| `{row['element']}` | `{row['file']}` | `{row['sha256']}` | "
                f"{row['status']} |\n"
            )
        report.write(
            "\n## 読取方式\n\n"
            "- 原Excelは変更していない。\n"
            "- 古いBIFFレコードとの互換性確保のため、LibreOffice Calcで一時XLSXを生成して読み取った。\n"
            "- 正本生成に使用した原Excelは、ZIP内部ファイルのSHA-256で識別できる。\n"
        )


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--year", type=int, required=True, choices=range(1953, 2000))
    parser.add_argument("--month", type=int, required=True, choices=range(1, 13))
    parser.add_argument("--source-directory", type=Path)
    parser.add_argument("--download-missing", action="store_true")
    return parser.parse_args()


def main() -> int:
    arguments = parse_arguments()
    manifest_path, manifest = load_manifest(arguments.year)
    configure_common_module(arguments.year, manifest_path)
    source_directory = arguments.source_directory or (
        REPOSITORY_ROOT / "build" / f"source-{arguments.year}"
    )
    source_paths, workbook_rows, archive_rows = prepare_source_files(
        manifest, source_directory, arguments.download_missing
    )
    extractions = extract_all_elements(source_paths, arguments.year, arguments.month)
    document = common.assemble_month_document(arguments.month, extractions)
    document["daily_summaries"] = extract_daily_summaries(
        source_paths, arguments.year, arguments.month
    )
    document["dataset"]["source_manifest"] = f"../../../metadata/sources/{manifest_path.name}"
    document["dataset"]["generator"] = GENERATOR_PATH
    document["notes"] = [
        f"1953–2007年一括ZIP内の{arguments.year}年Excelから生成した。",
        "原Excelは変更せず、ZIPと内部ExcelのSHA-256を検証した。",
        "古いBIFFレコードを読むため、一時XLSXへ変換した後に値を抽出した。",
        "1時から24時という原資料上の時刻表記をsource_hourへ保持する。",
        "日平均・日合計・最高・最低などは時間値と混在させずdaily_summariesへ収録した。",
    ]

    yaml_path = REPOSITORY_ROOT / "data" / STATION_ID / str(arguments.year) / f"{arguments.month:02d}.yaml"
    report_path = REPOSITORY_ROOT / "reports" / f"validation-{arguments.year}-{arguments.month:02d}.md"
    common.write_yaml(yaml_path, document)
    common.validate_round_trip(yaml_path, document)
    common.write_validation_report(
        report_path,
        arguments.month,
        yaml_path,
        document,
        extractions,
        workbook_rows,
    )
    append_historical_report_details(report_path, archive_rows)
    print(f"Generated: {yaml_path.relative_to(REPOSITORY_ROOT)}")
    print(f"Validated: {report_path.relative_to(REPOSITORY_ROOT)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, KeyError, TypeError, ValueError, subprocess.SubprocessError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
