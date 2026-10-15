#!/usr/bin/env python3
"""独立に読取用XLSXのセルと1954・1955年正本の観測値を全件照合する。"""

from __future__ import annotations

import argparse
import calendar
import json
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils.datetime import to_excel

from import_hitachi_city_hall_1954_1955 import (
    ROOT, SEA_LEVEL_SCHEDULES, prepare_sources,
)


FOUR = {6: 2, 9: 3, 14: 4, 22: 5}
DAYLIGHT = {hour: hour - 2 for hour in range(4, 21)}


def classify(value, *, direction=False, integer=False):
    if value is None:
        return None, "source_blank", None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None, "source_whitespace" if value else "source_blank", value or None
        if stripped == "-":
            return None, "source_dash", value
        if stripped == "***":
            return None, "source_triple_asterisk", value
        if stripped == "欠測":
            return None, "source_missing", value
        return (stripped.upper() if direction else stripped), None, (value if value != stripped else None)
    if isinstance(value, (int, float)):
        number = int(value) if integer and float(value).is_integer() else float(value)
        return number, None, None
    raise TypeError(f"Unclassified source value: {value!r}")


def nested_value(document, dotted_path):
    value = document
    for key in dotted_path.split("."):
        value = value[key]
    return value


def verify_daily_summaries(year, month, source_sheets, summaries):
    specifications = {
        "temperature.mean_c": ("temperature", 6, "number"),
        "temperature.maximum.value_c": ("temperature", 7, "number"),
        "temperature.maximum.time": ("temperature", 8, "time"),
        "temperature.minimum.value_c": ("temperature", 9, "number"),
        "temperature.minimum.time": ("temperature", 10, "time"),
        "humidity.mean_percent": ("humidity", 6, "number"),
        "humidity.minimum.value_percent": ("humidity", 7, "number"),
        "humidity.minimum.time": ("humidity", 8, "time"),
        "precipitation.total_mm": ("precipitation", 26, "number"),
        "precipitation.maximum_one_hour.value_mm": ("precipitation", 28, "number"),
        "precipitation.maximum_one_hour.time": ("precipitation", 29, "time"),
        "precipitation.maximum_ten_minutes.value_mm": ("precipitation", 31, "number"),
        "precipitation.maximum_ten_minutes.time": ("precipitation", 32, "time"),
        "station_pressure.mean_hpa": ("pressure", 6, "number"),
        "station_pressure.maximum.value_hpa": ("pressure", 7, "number"),
        "station_pressure.maximum.time": ("pressure", 8, "time"),
        "station_pressure.minimum.value_hpa": ("pressure", 9, "number"),
        "station_pressure.minimum.time": ("pressure", 10, "time"),
        "sunshine_duration.total_h": ("sunshine_duration", 19, "duration"),
        "wind.mean_speed_m_s": ("wind_speed", 26, "number"),
        "wind.maximum_10_minute.direction": ("wind_speed", 29, "text"),
        "wind.maximum_10_minute.speed_m_s": ("wind_speed", 30, "number"),
        "wind.maximum_10_minute.time": ("wind_speed", 31, "time"),
        "wind.maximum_instantaneous.direction": ("wind_speed", 33, "text"),
        "wind.maximum_instantaneous.speed_m_s": ("wind_speed", 34, "number"),
        "wind.maximum_instantaneous.time": ("wind_speed", 35, "time"),
    }
    compared = 0
    for day, summary in enumerate(summaries, start=1):
        row_numbers = {"temperature": day + 2, "humidity": day + 2,
                       "precipitation": day + 3, "pressure": day + 2,
                       "sunshine_duration": day + 2, "wind_speed": day + 3}
        for dotted_path, (sheet_name, column, value_kind) in specifications.items():
            raw = source_sheets[sheet_name].cell(row_numbers[sheet_name], column).value
            expected_flag = None
            expected_raw = None
            if value_kind == "number" or value_kind == "duration":
                expected, expected_flag, expected_raw = classify(raw)
                if isinstance(raw, timedelta) and value_kind == "duration":
                    expected, expected_flag, expected_raw = raw.total_seconds() / 3600, None, None
                elif isinstance(expected, str) and expected_flag is None:
                    expected, expected_flag, expected_raw = None, "unrecognized_source_value", raw
            elif value_kind == "time":
                if raw is None:
                    expected, expected_flag = None, "source_blank"
                elif isinstance(raw, time):
                    expected = raw.isoformat(timespec="minutes")
                elif isinstance(raw, datetime):
                    serial = float(to_excel(raw))
                    if serial >= 1:
                        expected = f"source_excel_serial:{serial:g}"
                        expected_raw = serial
                    else:
                        expected = raw.time().isoformat(timespec="minutes")
                elif isinstance(raw, timedelta):
                    seconds = int(raw.total_seconds())
                    if seconds >= 86400:
                        # Keep a date-serial-looking source cell explicit instead
                        # of silently treating it as the end of a day.
                        expected = f"source_excel_serial:{seconds / 86400:g}"
                        expected_raw = seconds / 86400
                    else:
                        hour, remainder = divmod(seconds, 3600)
                        minute, second = divmod(remainder, 60)
                        expected = f"{hour:02d}:{minute:02d}" if second == 0 else f"{hour:02d}:{minute:02d}:{second:02d}"
                else:
                    expected, expected_flag, expected_raw = classify(raw)
            else:
                expected, expected_flag, expected_raw = classify(raw)
                if isinstance(expected, (int, float)):
                    expected = str(expected)

            actual = nested_value(summary, dotted_path)
            actual_flag = summary.get("flags", {}).get(dotted_path)
            actual_raw = summary.get("raw_values", {}).get(dotted_path)
            if (actual, actual_flag, actual_raw) != (expected, expected_flag, expected_raw):
                raise AssertionError(
                    f"{year}-{month:02d}-{day:02d} summary {dotted_path}: "
                    f"actual={(actual, actual_flag, actual_raw)!r}, "
                    f"expected={(expected, expected_flag, expected_raw)!r}"
                )
            compared += 1
    return compared


def verify_year(year: int, source_paths: dict[str, Path]):
    canonical = {}
    for path in sorted((ROOT / "data/hitachi-city-hall" / str(year)).glob("*.yaml")):
        document = yaml.safe_load(path.read_text(encoding="utf-8"))
        for item in document["observations"]:
            canonical[(item["source_date"], item["source_hour"])] = item
    if len(canonical) != 365 * 24:
        raise AssertionError(f"{year}: expected 8760 observation-time records, got {len(canonical)}")

    books = {name: openpyxl.load_workbook(path, data_only=True)
             for name, path in source_paths.items()}
    compared = 0
    for month in range(1, 13):
        days = calendar.monthrange(year, month)[1]
        sheets = {name: book.worksheets[month - 1] for name, book in books.items()
                  if name != "weather_code"}
        weather = books["weather_code"]["天気"]
        month_document = yaml.safe_load(
            (ROOT / "data/hitachi-city-hall" / str(year) / f"{month:02d}.yaml")
            .read_text(encoding="utf-8")
        )
        month_summaries = month_document["daily_summaries"]
        precipitation_layout = month_document["dataset"]["source_table_layouts"]["precipitation_mm"]
        if precipitation_layout.get("status") != "no_hourly_values_in_month":
            raise AssertionError(f"{year}-{month:02d}: hourly precipitation layout status mismatch")
        sea_hours = SEA_LEVEL_SCHEDULES[year][month]
        for day in range(1, days + 1):
            source_date = f"{year}-{month:02d}-{day:02d}"
            for element, source_name, hours, row_number, columns, direction, integer in [
                ("temperature_c", "temperature", list(FOUR), day + 2, FOUR, False, False),
                ("relative_humidity_percent", "humidity", list(FOUR), day + 2, FOUR, False, False),
                ("station_pressure_hpa", "pressure", list(FOUR), day + 2, FOUR, False, False),
                ("sea_level_pressure_hpa", "pressure", sea_hours, day + 38,
                 {hour: index + 2 for index, hour in enumerate(sea_hours)}, False, False),
                ("sunshine_duration_h", "sunshine_duration", list(DAYLIGHT), day + 2,
                 DAYLIGHT, False, False),
                ("wind_speed_m_s", "wind_speed", list(range(1, 25)), day + 3,
                 {hour: hour + 1 for hour in range(1, 25)}, False, False),
                ("wind_direction", "wind_direction", list(range(1, 25)), day + 2,
                 {hour: hour + 1 for hour in range(1, 25)}, True, False),
            ]:
                for hour in hours:
                    source_value = sheets[source_name].cell(row_number, columns[hour]).value
                    expected, expected_flag, expected_raw = classify(source_value, direction=direction,
                                                                      integer=integer)
                    record = canonical[(source_date, hour)]
                    actual = record.get("values", {}).get(element, "__not_stored__")
                    actual_flag = record.get("flags", {}).get(element)
                    actual_raw = record.get("raw_values", {}).get(element)
                    if (actual, actual_flag, actual_raw) != (expected, expected_flag, expected_raw):
                        raise AssertionError(
                            f"{year}-{month:02d}-{day:02d} {hour}:00 {element}: "
                            f"actual={(actual, actual_flag, actual_raw)!r}, "
                            f"expected={(expected, expected_flag, expected_raw)!r}"
                        )
                    compared += 1

            # These workbooks provide a 1–24 precipitation grid, but every
            # hourly cell is blank; daily totals are stored separately.
            if any(sheets["precipitation"].cell(day + 3, hour + 1).value is not None
                   for hour in range(1, 25)):
                raise AssertionError(f"{source_date}: unexpected hourly precipitation source value")

            # The monthly weather code column is the source's numeric symbol column.
            weather_hour = 9
            source_value = weather.cell(day + 4, 19 + month).value
            expected, expected_flag, expected_raw = classify(source_value, integer=True)
            record = canonical[(source_date, weather_hour)]
            actual = record.get("values", {}).get("weather_code", "__not_stored__")
            actual_flag = record.get("flags", {}).get("weather_code")
            actual_raw = record.get("raw_values", {}).get("weather_code")
            if (actual, actual_flag, actual_raw) != (expected, expected_flag, expected_raw):
                raise AssertionError(f"{source_date} weather_code mismatch: {actual!r} != {expected!r}")
            compared += 1
        compared += verify_daily_summaries(year, month, sheets, month_summaries)
    return compared


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--year", type=int, choices=(1954, 1955), action="append")
    parser.add_argument("--source-directory", type=Path)
    parser.add_argument("--download-missing", action="store_true")
    args = parser.parse_args()
    years = args.year or [1954, 1955]
    results = {}
    for year in years:
        directory = args.source_directory or ROOT / "build" / f"source-{year}"
        _, source_paths, _ = prepare_sources(year, directory, args.download_missing)
        results[str(year)] = verify_year(year, source_paths)
        print(f"{year}: independently compared {results[str(year)]} source cells")
    result_path = ROOT / "build" / "verification-1954-1955.json"
    result_path.write_text(json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Result: {result_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
